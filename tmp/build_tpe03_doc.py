from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image

ROOT = Path(r"G:\UNIVERSIDAD\CUARTO AÑO\ISW\Ingenieria_Calidad_de_Software_4K1_2026")
OUT = ROOT / r"04_Trabajos_Practicos\Evaluables\TPE03\G6_TPE03_Documentacion_TDD.docx"
REF_IMG = ROOT / r"tmp\reference_pdf\page1-300.png"
LOGO = ROOT / r"tmp\reference_pdf\utn_logo.png"

BLUE = "2F72B7"
DARK = "343642"
LIGHT_BLUE = "EAF2FA"
PALE = "F6F8FA"
GRID = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="8"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name="Arial", size=10.5, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(p, before=0, after=6, line=1.15, keep_next=False):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep_next


def add_text(doc, text, bold_prefix=None, after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    style_paragraph(p, after=after)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullets(doc, items, level=0, compact=False):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        style_paragraph(p, after=2 if compact else 4, line=1.08)
        r = p.add_run(item)
        set_font(r)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        style_paragraph(p, after=4, line=1.1)
        r = p.add_run(item)
        set_font(r)


def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, label in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = Inches(widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, after=0, line=1.0)
        r = p.add_run(label)
        set_font(r, size=font_size, bold=True, color="FFFFFF")
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            if row_idx % 2:
                set_cell_shading(cell, PALE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Inches(widths[i])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(p, after=0, line=1.05)
            r = p.add_run(str(value))
            set_font(r, size=font_size)
        tr_pr = table.rows[-1]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    return p


def page_break(doc):
    doc.add_page_break()


def add_header(section):
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    tabs = p.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(6.6))
    r = p.add_run("Cátedra: ")
    set_font(r, size=9.5, bold=True, color="000000")
    r = p.add_run("Ingeniería y Calidad de Software")
    set_font(r, size=9.5, color="000000")
    r = p.add_run("\t2026-4K1-Grupo 6-TP 3")
    set_font(r, size=9.5, bold=True, color="000000")
    p_pr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "22")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "2C61B5")
    pbdr.append(bottom)
    p_pr.append(pbdr)


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run = p.add_run()
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_font(run, size=8.5, color="666666")


def build_logo():
    img = Image.open(REF_IMG).convert("RGBA")
    # Coordenadas del isotipo UTN en la portada renderizada a 300 dpi.
    crop = img.crop((900, 1210, 1480, 1815))
    px = crop.load()
    for y in range(crop.height):
        for x in range(crop.width):
            r, g, b, a = px[x, y]
            if r > 245 and g > 245 and b > 245:
                px[x, y] = (255, 255, 255, 0)
    bbox = crop.getbbox()
    if bbox:
        crop = crop.crop(bbox)
    crop.save(LOGO)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("000000")

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("000000")

    for level, size in ((1, 17), (2, 13.5), (3, 11.5)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = level != 1
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.space_before = Pt(8 if level == 1 else 6)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    code = doc.styles.add_style("Código", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor.from_string("222222")
    code.paragraph_format.left_indent = Inches(0.25)
    code.paragraph_format.right_indent = Inches(0.25)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.0


def add_code(doc, text):
    p = doc.add_paragraph(style="Código")
    set_cell_shading_like_paragraph(p, "F3F5F7")
    for line in text.splitlines():
        r = p.add_run(line + "\n")
        set_font(r, name="Consolas", size=8.5)


def set_cell_shading_like_paragraph(p, fill):
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def build_doc():
    build_logo()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.62)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    sec.header_distance = Inches(0.25)
    sec.footer_distance = Inches(0.28)
    add_header(sec)
    add_footer(sec)
    configure_styles(doc)

    # Portada
    for _ in range(3):
        doc.add_paragraph()
    for text, size in [
        ("Universidad Tecnológica Nacional", 22),
        ("Facultad Regional Córdoba", 17),
        ("Ingeniería en Sistemas de Información", 16),
        ("Ingeniería y Calidad de Software", 13),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, after=18 if size > 16 else 10)
        r = p.add_run(text)
        set_font(r, size=size, bold=True, color=DARK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(LOGO), width=Inches(1.65))
    style_paragraph(p, after=8)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TEMA: Desarrollo Conducido por Pruebas")
    set_font(r, size=13, bold=True, color=DARK)
    style_paragraph(p, after=10)
    cover_lines = [
        "Unidad Nro. 4",
        "Trabajo Práctico Nro. 3",
        "Año: 2026",
        "Docentes:\n   •  Ing. Judith Meles\n   •  Ing. Valeria Abdala\n   •  Salvador Barbera",
        "Curso: 4K1\nGrupo Nro: 6",
    ]
    for idx, text in enumerate(cover_lines):
        p = doc.add_paragraph()
        style_paragraph(p, after=7)
        p.paragraph_format.left_indent = Inches(0.05)
        for j, line in enumerate(text.splitlines()):
            if j:
                p.add_run("\n")
            r = p.add_run(line)
            set_font(r, size=11.5, bold=(j == 0 and idx != 3) or (idx == 3 and j == 0), color=DARK)

    # Integrantes
    page_break(doc)
    add_heading(doc, "Integrantes", 1)
    integrantes = [
        ("Juan Ignacio Arrigoni", "407069", "juani.arrigoni@gmail.com"),
        ("Lautaro Francisco Castillo", "400998", "castillolautarof@gmail.com"),
        ("Francisco Dalmazzo", "400193", "fgdalmazzo@gmail.com"),
        ("Ezequiel Deza", "401006", "dezaezequiel5@gmail.com"),
        ("Gina Eusebi", "400202", "ginaeusebi9@gmail.com"),
        ("Juan Ignacio Fanfliet", "400879", "juanignaciofanfliet@gmail.com"),
        ("Olivia García Osella", "400032", "oliviaosella@gmail.com"),
        ("Delfina Guerci", "401782", "delfinaguerci@gmail.com"),
        ("Ignacio Facundo Lescano Barreto", "400523", "facuiglescano14@gmail.com"),
        ("Ignacio Martín", "404677", "ignamartin112@gmail.com"),
        ("Elena Adriana Moyano Matlega", "401342", "moyanoelena7@gmail.com"),
        ("Catalina Pérez", "402462", "cataaperezz04@gmail.com"),
        ("Diego Matías Petitto", "400258", "dmpetitto2004@gmail.com"),
        ("Ignacio Samocachan", "400621", "ignasamo2@gmail.com"),
        ("Pedro Torcivia", "400091", "pedritorcivia@gmail.com"),
    ]
    add_table(doc, ["Nombre y apellido", "Legajo", "Email"], integrantes, [2.45, 1.05, 3.35], font_size=9.1)

    # Índice
    page_break(doc)
    add_heading(doc, "Índice", 1)
    toc = [
        ("1. Introducción", "4"), ("2. Consigna y alcance", "4"),
        ("3. Solución implementada", "5"), ("4. Aplicación del enfoque TDD", "6"),
        ("5. Evidencia por subgrupo", "7"), ("5.1 Consulta de actividades y turnos", "7"),
        ("5.2 Términos y aceptación", "8"), ("5.3 Validación de participantes", "9"),
        ("5.4 Confirmación correo y QR", "10"), ("5.5 Reservas y cupos", "11"),
        ("6. Integración de los módulos", "12"), ("7. Pruebas automatizadas", "13"),
        ("8. Decisiones de diseño", "14"), ("9. Ejecución y configuración", "14"),
        ("10. Conclusión", "15"), ("11. Glosario", "15"),
    ]
    for label, page in toc:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28 if label.startswith("5.") and len(label.split(".")[0:2]) > 1 else 0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.55))
        r = p.add_run(label)
        set_font(r, size=10.5, bold=label.split(".")[0].isdigit() and label.count(".") == 1)
        r = p.add_run(f"\t{page}")
        set_font(r, size=10.5, bold=True)

    # Página 4
    page_break(doc)
    add_heading(doc, "1. Introducción", 1)
    add_text(doc, "El presente informe documenta la implementación de la User Story Inscribirme a actividad de EcoHarmony Park mediante Desarrollo Conducido por Pruebas. El trabajo se realizó en Python con FastAPI para el backend y React con TypeScript para la interfaz. La construcción del comportamiento de negocio se inició desde pruebas automatizadas y quedó registrada en ramas y commits que identifican las etapas Red, Green y Refactor.")
    add_text(doc, "El resultado integra consulta de actividades y turnos, aceptación de términos, validación de participantes, gestión de reservas y cupos, generación del comprobante QR y envío de la confirmación por correo. La suite integrada alcanza 171 casos automatizados y utiliza dependencias falsas o temporales para evitar envíos reales y aislar la persistencia durante las pruebas.")
    add_heading(doc, "2. Consigna y alcance", 1)
    add_text(doc, "La consigna solicitó implementar una funcionalidad mediante TDD y aplicar el ciclo Red Green Refactor para construir el código desde las pruebas. Por tratarse del Grupo 6, se seleccionó la User Story Inscribirme a actividad.")
    add_heading(doc, "2.1 Objetivo y propósito", 2)
    add_text(doc, "El objetivo fue diseñar y construir una solución funcional a partir de pruebas automatizadas, incorporando criterios de calidad desde el inicio. El propósito fue ejercitar la disciplina TDD y utilizar cada prueba como especificación ejecutable del comportamiento esperado.")
    add_heading(doc, "2.2 Entradas y entregables", 2)
    add_bullets(doc, [
        "Conceptos teóricos y bibliografía de la Unidad 4.",
        "Enunciado de EcoHarmony Park, User Stories del Trabajo Práctico 2 y reglas de estilo del grupo.",
        "Código funcional de la inscripción a actividades y pruebas automatizadas de sus casos.",
        "Evidencia del ciclo TDD y justificación de las decisiones de diseño e integración.",
    ], compact=True)

    # Página 5
    page_break(doc)
    add_heading(doc, "3. Solución implementada", 1)
    add_text(doc, "La aplicación permite seleccionar una actividad, consultar fechas y turnos disponibles, cargar participantes, aceptar las condiciones vigentes y confirmar la reserva. El backend vuelve a validar las reglas en el registro definitivo, descuenta cupos de forma atómica, persiste la reserva y genera un QR legible. El correo se procesa después de confirmar la reserva; si falla, la inscripción conserva su validez y puede reintentarse sin descontar cupos nuevamente.")
    add_heading(doc, "3.1 Arquitectura", 2)
    architecture = [
        ("Frontend React", "Presenta el flujo y consume exclusivamente datos y reglas publicados por la API."),
        ("API FastAPI", "Expone actividades, calendario, horarios, términos, validación y registro de inscripciones."),
        ("Servicios de dominio", "Aplican reglas de participantes, aceptación, reservas, QR y notificaciones."),
        ("Persistencia SQLite", "Conserva reservas y controla cupos por actividad, fecha y horario."),
        ("Servicios externos", "SMTP entrega el correo; la librería qrcode genera el comprobante PNG."),
    ]
    add_table(doc, ["Componente", "Responsabilidad"], architecture, [1.65, 5.2], font_size=9.4)
    add_heading(doc, "3.2 Flujo funcional", 2)
    add_numbered(doc, [
        "Consultar el catálogo oficial de actividades.",
        "Consultar el calendario y los turnos reservables con sus cupos.",
        "Cargar los participantes y validar sus datos mediante el backend.",
        "Obtener y aceptar la versión vigente de los términos de la actividad.",
        "Registrar la inscripción; el servidor revalida apertura, cupos, participantes y términos.",
        "Mostrar el QR y el estado del correo. Si el envío falla, permitir el reintento de la notificación.",
    ])
    add_heading(doc, "3.3 Contratos principales", 2)
    add_table(doc, ["Método y ruta", "Uso"], [
        ("GET /actividades", "Catálogo, capacidad, edad mínima, talle y talles disponibles."),
        ("GET /actividades/{actividad}/calendario", "Ventana, cierres, fechas, horarios, identificadores y cupos."),
        ("POST /inscripciones/validar", "Validación informativa de participantes."),
        ("GET /actividades/{actividad}/terminos", "Texto y versión de las condiciones vigentes."),
        ("POST /inscripciones", "Confirmación integral y persistencia de la reserva."),
    ], [2.55, 4.3], font_size=8.7)

    # Página 6
    page_break(doc)
    add_heading(doc, "4. Aplicación del enfoque TDD", 1)
    add_text(doc, "El trabajo se dividió por responsabilidades para permitir ciclos pequeños. Cada subgrupo desarrolló una porción del comportamiento en una rama independiente y registró en los mensajes de commit la fase del ciclo aplicada. Luego se integraron los contratos y se agregaron pruebas que cubren la colaboración entre módulos.")
    add_heading(doc, "4.1 Ciclo Red Green Refactor", 2)
    add_table(doc, ["Fase", "Aplicación en el proyecto"], [
        ("Red", "Se escribió una prueba que expresaba el comportamiento faltante y se verificó el fallo esperado."),
        ("Green", "Se agregó la implementación mínima necesaria para que la prueba y la suite existente pasaran."),
        ("Refactor", "Se reorganizó el diseño, se extrajeron helpers y se eliminaron duplicaciones sin cambiar el comportamiento."),
    ], [1.15, 5.7], font_size=9.4)
    add_heading(doc, "4.2 Ramas de trabajo", 2)
    add_table(doc, ["Rama", "Responsabilidad"], [
        ("tdd/consultar-actividad", "Catálogo, turnos, cierres, anticipación y cupos visibles."),
        ("tdd/terminos-confirmacion", "Términos por actividad y aceptación obligatoria."),
        ("tdd/cargar-participante", "Cantidad, nombre, DNI, edad, talle y duplicados."),
        ("tdd/envio-confirmacion-inscripcion", "Resumen, correo, QR e inyección de dependencias."),
        ("gestion-reservas-cupos", "Reserva, persistencia, cupos, solapamiento y concurrencia."),
        ("tdd/integracion/integracion-grupos-1-a-5", "Contrato HTTP y orquestación de los cinco módulos."),
        ("tdd/frontend", "Interfaz conectada al contrato integrado del backend."),
    ], [2.65, 4.2], font_size=8.7)
    add_text(doc, "La evidencia incluye un revert explícito del commit que incorporaba SMTP sin una prueba previa. La funcionalidad se reimplementó después con su ciclo Red y Green, lo que muestra que el proceso TDD se utilizó también como criterio de control del historial.")

    # Página 7 Grupo 1
    page_break(doc)
    add_heading(doc, "5. Evidencia por subgrupo", 1)
    add_heading(doc, "5.1 Consulta de actividades y turnos", 2)
    add_text(doc, "El primer módulo construyó el catálogo y la consulta de turnos. Las pruebas definieron los horarios entre las 09:00 y las 17:00, la independencia de cupos por turno, el estado sin cupos y los cierres de los lunes, el 1 de enero y el 25 de diciembre. También se agregó una ventana adoptada por el backend: el turno se muestra cuando faltan al menos 15 minutos y como máximo 24 horas.")
    add_heading(doc, "Evidencia de commits", 3)
    add_table(doc, ["Fase", "Commit", "Evidencia"], [
        ("Red", "7136bf3", "Prueba que consulta turnos un lunes y espera una lista vacía."),
        ("Green", "d0f0d7a", "Implementación del cierre semanal."),
        ("Red", "2101455", "Prueba de turno completo sin disponibilidad."),
        ("Green", "26cdd7d", "Estado y motivo para el turno lleno."),
        ("Refactor", "5dcebf4", "Reorganización del cálculo independiente de cupos."),
        ("Red", "51a82e1", "Caso con menos de 15 minutos de anticipación."),
        ("Green", "580a29f", "Filtro por anticipación mínima."),
        ("Green", "72e2669", "Filtro para turnos a más de 24 horas."),
    ], [0.8, 1.0, 5.05], font_size=8.6)
    add_heading(doc, "Decisiones de diseño", 3)
    add_bullets(doc, [
        "Representar cada turno con actividad, fecha, hora, cupos disponibles, disponibilidad y motivo.",
        "Consultar la ocupación mediante una dependencia para aislar el cálculo durante los tests.",
        "Centralizar los cierres y los límites de anticipación en el backend para que el frontend no replique reglas.",
    ])

    # Página 8 Grupo 2
    page_break(doc)
    add_heading(doc, "5.2 Términos y aceptación", 2)
    add_text(doc, "El segundo módulo estableció que cada inscripción debe aceptar términos correspondientes a la misma actividad. Los casos parametrizados cubren las cuatro actividades, la ausencia de aceptación, la aceptación cruzada y la prioridad de errores. Los textos se separaron de la lógica para permitir que cambien sin alterar la validación.")
    add_heading(doc, "Evidencia de commits", 3)
    add_table(doc, ["Fase", "Commit", "Evidencia"], [
        ("Red", "2ffade1", "Prueba de rechazo cuando los términos no fueron aceptados."),
        ("Green", "65665f6", "Excepción de términos no aceptados."),
        ("Red", "feb0270", "Prueba para aceptación perteneciente a otra actividad."),
        ("Green", "65e00e4", "Validación de correspondencia entre actividad y aceptación."),
        ("Refactor", "64acfe9", "Unificación de combinaciones de actividades en tests parametrizados."),
        ("Red", "1d91ec3", "Prueba inicial de obtención de términos con NotImplementedError."),
        ("Green", "2a1d74b", "Texto mínimo para satisfacer el contrato."),
        ("Refactor", "e1d496b", "Separación del texto ficticio y la función de consulta."),
    ], [0.8, 1.0, 5.05], font_size=8.6)
    add_heading(doc, "Decisiones de diseño", 3)
    add_bullets(doc, [
        "Modelar la aceptación como un objeto con actividad, indicador de aceptación y versión.",
        "Validar la aceptación antes de modificar cupos o persistir una reserva.",
        "Enviar la versión mostrada al usuario y rechazar una confirmación si los términos cambiaron.",
    ])

    # Página 9 Grupo 3
    page_break(doc)
    add_heading(doc, "5.3 Validación de participantes", 2)
    add_text(doc, "El tercer módulo definió las reglas de los visitantes: debe existir al menos un participante, la cantidad declarada debe coincidir con la lista, el nombre debe contener partes alfabéticas, el DNI debe ser numérico y positivo, la edad no puede ser negativa y debe respetar el mínimo de la actividad, y el talle debe pertenecer al catálogo cuando sea obligatorio. También se rechazan DNI repetidos dentro de la misma solicitud.")
    add_heading(doc, "Evidencia de commits", 3)
    add_table(doc, ["Fase", "Commit", "Evidencia"], [
        ("Red", "bc1f855", "Prueba inicial de existencia de las validaciones."),
        ("Green", "63bda28", "Implementación mínima del contrato."),
        ("Red", "8e1a4c6", "Prueba de nombre válido."),
        ("Green", "686763f", "Validación del nombre del participante."),
        ("Red", "0d804a1", "Prueba de talle obligatorio."),
        ("Green", "1254d75", "Validación de talle según actividad."),
        ("Red", "bccb5b4", "Prueba de DNI repetidos."),
        ("Green", "ebeb551", "Rechazo de duplicados."),
        ("Refactor", "dcd2e4f", "Helper de construcción de participantes para reducir duplicación."),
    ], [0.8, 1.0, 5.05], font_size=8.5)
    add_heading(doc, "Decisiones de diseño", 3)
    add_bullets(doc, [
        "Mantener una validación individual reutilizable y una validación del conjunto.",
        "Usar el enum Actividad como fuente de cupo, edad mínima y obligatoriedad de talle.",
        "No inventar límites de DNI o edad que no estén definidos por el contrato probado.",
    ])

    # Página 10 Grupo 4
    page_break(doc)
    add_heading(doc, "5.4 Confirmación correo y QR", 2)
    add_text(doc, "El cuarto módulo prepara un resumen legible, genera un QR PNG y envía el comprobante por correo. El diseño utiliza inyección de dependencias: CorreoFalso y SMTPFalso permiten verificar destinatario, contenido, adjunto y orden de las operaciones sin conectarse a un servidor real. Las credenciales se leen desde variables de entorno y no forman parte del repositorio.")
    add_heading(doc, "Evidencia de commits", 3)
    add_table(doc, ["Fase", "Commit", "Evidencia"], [
        ("Red", "e336009", "Prueba inicial del resumen de confirmación."),
        ("Green", "d867288", "Implementación mínima; no se refactorizó porque ya era suficiente."),
        ("Red", "4dd9ab7", "Prueba de envío mediante CorreoFalso."),
        ("Green", "7a6217c", "Invocación mínima del emisor inyectado."),
        ("Red", "c69376c", "Prueba que exige el QR dentro del correo."),
        ("Green", "73f1b76", "Integración de generar_qr en el servicio."),
        ("Refactor", "c2e39ac", "Reorganización de imports y mejora del resumen."),
        ("Control", "027a242", "Revert del código SMTP agregado sin test previo."),
    ], [0.8, 1.0, 5.05], font_size=8.5)
    add_heading(doc, "Decisiones de diseño", 3)
    add_bullets(doc, [
        "Separar la generación de QR del envío de correo para probarlos en forma independiente.",
        "Usar smtplib y EmailMessage, con configuración SMTP externa mediante .env.",
        "Confirmar la reserva aunque el correo falle y registrar el estado para permitir reintentos.",
    ])

    # Página 11 Grupo 5
    page_break(doc)
    add_heading(doc, "5.5 Reservas y cupos", 2)
    add_text(doc, "El quinto módulo administra la creación de reservas y la disponibilidad. Las pruebas comenzaron con capacidades iniciales, continuaron con descuento e independencia de cupos y cubrieron errores sin efectos colaterales, persistencia, DNI superpuesto y dos solicitudes concurrentes por los últimos lugares. La versión integrada reemplazó la persistencia en memoria inicial por SQLite sin perder el contrato probado.")
    add_heading(doc, "Evidencia de commits", 3)
    add_table(doc, ["Fase", "Commit", "Evidencia"], [
        ("Red", "b096643", "Prueba de descuento de cupos en una reserva válida."),
        ("Green", "01e7125", "Descuento mínimo de la cantidad solicitada."),
        ("Refactor", "bba8d41", "Mejora de la estructura del control de cupos."),
        ("Red", "3e0e37d", "Prueba de rechazo sin persistencia ni modificación de cupos."),
        ("Green", "d09cb7f", "Validaciones previas a cualquier mutación."),
        ("Red", "cb85c56", "Dos solicitudes simultáneas compiten por los últimos lugares."),
        ("Green", "3bceeb5", "Exclusión mutua para impedir la sobreventa."),
        ("Refactor", "7c41d5c", "Separación de lecturas internas para evitar bloqueos."),
    ], [0.8, 1.0, 5.05], font_size=8.5)
    add_heading(doc, "Decisiones de diseño", 3)
    add_bullets(doc, [
        "Ejecutar todas las validaciones antes de persistir o descontar cupos.",
        "Usar una operación atómica para comprobar y consumir disponibilidad.",
        "Separar cupos y DNI por actividad, fecha e identificador de turno.",
    ])

    # Página 12 Integración
    page_break(doc)
    add_heading(doc, "6. Integración de los módulos", 1)
    add_text(doc, "La integración requirió un contrato común para actividad, fecha, horario, participantes, aceptación, correo y respuesta. Los endpoints conservan las reglas de los subgrupos y las aplican nuevamente al confirmar. Esta decisión evita que una consulta previa de disponibilidad se confunda con una reserva efectiva.")
    add_heading(doc, "6.1 Decisiones de integración", 2)
    add_table(doc, ["Problema", "Resolución integrada"], [
        ("Modelos diferentes entre ramas", "Se consolidaron enums y modelos Pydantic en schemas.py."),
        ("Cupos consultados antes de confirmar", "El POST definitivo repite la validación y descuenta en una transacción."),
        ("Persistencia inicialmente en memoria", "SQLite conserva reservas y evita sobreventa mediante transacciones."),
        ("QR corto en la rama inicial", "El comprobante final contiene identificador y datos legibles de la reserva."),
        ("Fallo de SMTP", "La reserva permanece confirmada; se informa correo_estado y se permite reintentar."),
        ("Términos cambiantes", "La versión aceptada se compara con la vigente antes de registrar."),
    ], [2.25, 4.6], font_size=9.0)
    add_heading(doc, "6.2 Frontend basado en el backend", 2)
    add_text(doc, "La interfaz no conserva un catálogo paralelo de reglas. Obtiene del backend actividades, cupos, edades mínimas, talles, cierres, ventana de anticipación, horarios e identificadores. La validación de participantes se ejecuta mediante POST /inscripciones/validar y la confirmación vuelve a pasar por el backend. En el frontend permanecen únicamente el estado temporal del formulario y los recursos visuales asociados a cada actividad.")
    add_heading(doc, "6.3 Commits de integración", 2)
    add_table(doc, ["Commit", "Resultado"], [
        ("5140e41", "Integración de actividades y participantes con términos, reservas y envío."),
        ("91d4d5f", "Cierre de QR, SMTP, términos y persistencia."),
        ("3749398", "Comprobante QR legible y datos organizados."),
        ("442064a", "Backend integrado con el frontend."),
        ("c0ad04c", "Backend convertido en fuente de reglas de la interfaz."),
    ], [1.2, 5.65], font_size=9.0)

    # Página 13 Tests
    page_break(doc)
    add_heading(doc, "7. Pruebas automatizadas", 1)
    add_text(doc, "La verificación final recoge 171 casos en 15 archivos. Los conteos incluyen parametrizaciones, por lo que superan la cantidad de funciones cuyo nombre comienza con test. La suite utiliza relojes controlados, bases temporales y servicios falsos para mantener resultados repetibles.")
    tests = [
        ("test_consulta_actividades.py", "16", "Catálogo, turnos, cierres, anticipación y cupos"),
        ("test_endpoints.py", "7", "Contratos HTTP iniciales"),
        ("test_terminos.py", "12", "Textos y estabilidad de términos"),
        ("test_aceptacion.py", "32", "Aceptación y prioridad de errores"),
        ("test_participantes.py", "12", "Nombre, DNI, edad, talle y cantidad"),
        ("test_notificaciones.py", "8", "Resumen, destinatario e inyección"),
        ("test_generacion_qr.py", "3", "PNG, contenido y unicidad"),
        ("test_comprobante_legible.py", "1", "Coherencia del comprobante"),
        ("test_correo_smtp.py", "1", "SMTP con servidor falso"),
        ("test_reservas.py", "12", "Cupos, persistencia y concurrencia"),
        ("test_integracion_inscripcion.py", "3", "Contrato compartido"),
        ("test_integracion.py", "8", "Orquestación de servicios"),
        ("test_integracion_completa.py", "20", "Recorrido HTTP y reglas cruzadas"),
        ("test_api_endpoints.py", "21", "API, calendario, cierres y CORS"),
        ("test_cierre_backend.py", "15", "SQLite, QR completo y reintentos"),
    ]
    add_table(doc, ["Archivo", "Casos", "Alcance"], tests, [2.25, 0.65, 3.95], font_size=8.2)
    add_heading(doc, "7.1 Resultado final", 2)
    add_code(doc, "........................................................................ [ 42%]\n........................................................................ [ 84%]\n...........................                                              [100%]\n171 passed")

    # Página 14 Decisiones y ejecución
    page_break(doc)
    add_heading(doc, "8. Decisiones de diseño", 1)
    decisions = [
        ("D1", "Reglas en el backend", "La API es la única fuente de reglas; la interfaz solo representa su respuesta."),
        ("D2", "Validación repetida", "La consulta previa mejora la interacción, pero el POST definitivo protege la integridad."),
        ("D3", "Persistencia SQLite", "Permite recuperar reservas y realizar el descuento dentro de una transacción."),
        ("D4", "Dependencias inyectables", "Correo, reloj y persistencia se sustituyen en tests sin efectos externos."),
        ("D5", "Reserva independiente del correo", "Una caída de SMTP no invalida ni duplica la reserva."),
        ("D6", "Contratos tipados", "Pydantic valida solicitudes y respuestas; los errores HTTP expresan la causa."),
    ]
    add_table(doc, ["ID", "Decisión", "Justificación"], decisions, [0.55, 1.75, 4.55], font_size=8.8)
    add_heading(doc, "9. Ejecución y configuración", 1)
    add_heading(doc, "9.1 Backend", 2)
    add_code(doc, "cd 04_Trabajos_Practicos/Evaluables/TPE03/G6_COD_InscripcionActividad/backend\n.\\.venv\\Scripts\\python.exe -m uvicorn app.main:app --reload")
    add_heading(doc, "9.2 Frontend", 2)
    add_code(doc, "cd 04_Trabajos_Practicos/Evaluables/TPE03/G6_COD_InscripcionActividad/frontend\nnpm ci\nnpm run dev")
    add_heading(doc, "9.3 Variables de correo", 2)
    add_text(doc, "La configuración SMTP se guarda en backend/.env y se toma de backend/.env.example. SMTP_HOST, SMTP_USER y SMTP_PASSWORD son obligatorias. La contraseña debe ser una contraseña de aplicación y el archivo .env permanece excluido del control de versiones.")
    add_heading(doc, "9.4 Ejecución de pruebas", 2)
    add_code(doc, "cd backend\n.\\.venv\\Scripts\\python.exe -m pytest -q\n\ncd ../frontend\nnpm run build\nnpm run lint")

    # Página 15 Conclusión y glosario
    page_break(doc)
    add_heading(doc, "10. Conclusión", 1)
    add_text(doc, "La User Story Inscribirme a actividad quedó implementada como un flujo completo y verificable. Los ciclos registrados muestran que las reglas se incorporaron a partir de pruebas pequeñas y que el diseño se ajustó sin perder comportamiento. La integración añadió pruebas de contrato y de recorrido para comprobar que las piezas desarrolladas por separado funcionaran juntas.")
    add_text(doc, "La solución final conserva las reglas en el backend, protege los cupos mediante persistencia transaccional, desacopla el envío de correo y expone al frontend un contrato único. La suite de 171 casos permite repetir la verificación y detectar regresiones en actividades, participantes, términos, reservas, comprobantes y endpoints.")
    add_heading(doc, "11. Glosario", 1)
    glossary = [
        ("TDD", "Desarrollo Conducido por Pruebas; práctica en la que las pruebas guían la implementación."),
        ("Red", "Etapa en la que una prueba nueva falla por la ausencia del comportamiento esperado."),
        ("Green", "Etapa en la que se agrega el código mínimo para que la prueba pase."),
        ("Refactor", "Mejora del diseño que conserva el comportamiento verificado por la suite."),
        ("Triangulación", "Uso de casos adicionales para evitar una implementación demasiado específica."),
        ("Doble de prueba", "Objeto controlado que sustituye una dependencia real, como CorreoFalso o SMTPFalso."),
        ("Prueba de integración", "Caso que verifica la colaboración entre módulos o capas."),
        ("Regresión", "Fallo que reaparece luego de un cambio y que una prueba automatizada permite detectar."),
        ("Atomicidad", "Propiedad por la cual una operación se completa en su totalidad o no modifica el estado."),
        ("Contrato HTTP", "Estructura acordada para solicitudes, respuestas, estados y errores de una API."),
    ]
    add_table(doc, ["Término", "Definición"], glossary, [1.35, 5.5], font_size=9.1)
    add_heading(doc, "11.1 Repositorio", 2)
    add_text(doc, "https://github.com/oliviaosella/G6_Ingenieria_Calidad_de_Software_4K1_2026.git", align=WD_ALIGN_PARAGRAPH.LEFT)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
